from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor

OUT = '四篇生成模型组会汇报_精简版.pptx'
NAVY = RGBColor(26, 50, 82)
BLUE = RGBColor(48, 105, 170)
LIGHT_BLUE = RGBColor(232, 242, 252)
CYAN = RGBColor(64, 150, 170)
ORANGE = RGBColor(222, 132, 55)
LIGHT_ORANGE = RGBColor(252, 241, 229)
GREEN = RGBColor(67, 132, 96)
LIGHT_GREEN = RGBColor(232, 245, 235)
RED = RGBColor(181, 79, 75)
LIGHT_RED = RGBColor(251, 235, 234)
GRAY = RGBColor(92, 103, 115)
MID = RGBColor(187, 197, 208)
LIGHT = RGBColor(246, 248, 250)
BLACK = RGBColor(27, 31, 36)
WHITE = RGBColor(255, 255, 255)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

def rect(slide, x, y, w, h, fill=WHITE, line=None, radius=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    s.line.color.rgb = line if line else fill
    s.line.width = Pt(0.8)
    if radius: s.adjustments[0] = 0.08
    return s

def line(slide, x1, y1, x2, y2, color=MID, width=1.4, arrow=False):
    s = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    s.line.color.rgb = color; s.line.width = Pt(width)
    if arrow: s.line.end_arrowhead = True
    return s

def text(slide, x, y, w, h, value, size=18, color=BLACK, bold=False,
         align=PP_ALIGN.LEFT, font='Microsoft YaHei', valign=MSO_ANCHOR.TOP,
         margin=0.04, italic=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = Inches(margin); tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin); tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]; p.alignment = align
    run = p.add_run(); run.text = value
    run.font.name = font; run.font.size = Pt(size); run.font.bold = bold
    run.font.italic = italic; run.font.color.rgb = color
    return box

def rich_text(slide, x, y, w, h, lines, size=16, color=BLACK, bullet=False,
              font='Microsoft YaHei', spacing=4):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = Inches(0.03); tf.margin_right = Inches(0.03)
    tf.margin_top = Inches(0.02); tf.margin_bottom = Inches(0.02)
    for i, item in enumerate(lines):
        if isinstance(item, tuple): val, col, b = item
        else: val, col, b = item, color, False
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ('• ' if bullet else '') + val
        p.font.name = font; p.font.size = Pt(size); p.font.color.rgb = col
        p.font.bold = b; p.space_after = Pt(spacing)
    return box

def footer(slide, n, source='四篇论文阅读笔记'):
    line(slide, 0.55, 7.08, 12.78, 7.08, MID, 0.7)
    text(slide, 0.6, 7.12, 8, 0.2, source, 8, GRAY)
    text(slide, 12.1, 7.1, 0.55, 0.24, f'{n:02d}', 9, GRAY, align=PP_ALIGN.RIGHT)

def title_bar(slide, kicker, title, subtitle=None, accent=BLUE):
    text(slide, 0.62, 0.34, 11.8, 0.26, kicker.upper(), 9, accent, True)
    text(slide, 0.6, 0.63, 12.1, 0.55, title, 27, NAVY, True)
    if subtitle: text(slide, 0.62, 1.19, 11.8, 0.34, subtitle, 11, GRAY)
    line(slide, 0.62, 1.62, 12.72, 1.62, MID, 1.0)

def pill(slide, x, y, w, label, fill, color=BLACK):
    rect(slide, x, y, w, 0.34, fill, fill, True)
    text(slide, x+0.04, y+0.03, w-0.08, 0.25, label, 10, color, True, PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)

def bullet_box(slide, x, y, w, h, heading, items, accent=BLUE, fill=LIGHT):
    rect(slide, x, y, w, h, fill, fill, True)
    text(slide, x+0.18, y+0.14, w-0.36, 0.3, heading, 15, accent, True)
    rich_text(slide, x+0.18, y+0.55, w-0.36, h-0.68, items, 13, BLACK, bullet=True, spacing=7)

def add_circle(slide, x, y, d, fill, label, label_color=WHITE, size=17):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    s.fill.solid(); s.fill.fore_color.rgb = fill; s.line.color.rgb = fill
    text(slide, x, y+0.01, d, d-0.02, label, size, label_color, True, PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    return s

# 1. Cover
slide = prs.slides.add_slide(blank)
rect(slide, 0, 0, 13.333, 7.5, WHITE, WHITE)
rect(slide, 0, 0, 0.18, 7.5, BLUE, BLUE)
text(slide, 0.78, 0.78, 4.2, 0.28, '组会汇报 / PAPER READING', 11, BLUE, True)
text(slide, 0.78, 1.38, 11.8, 0.74, '四篇生成模型：从扩散到漂移', 31, NAVY, True)
text(slide, 0.8, 2.22, 10.5, 0.42, '用一条主线理解四种“从噪声到数据”的建模方式', 17, GRAY)
line(slide, 0.82, 3.24, 12.25, 3.24, MID, 1.2)
models = [('DDPM', '离散链', BLUE), ('Score-SDE', '连续 SDE', CYAN), ('Flow Matching', '向量场', ORANGE), ('Drifting', '训练时漂移', GREEN)]
xs = [0.95, 3.95, 6.95, 9.95]
for i, (name, sub, col) in enumerate(models):
    add_circle(slide, xs[i], 2.98, 0.52, col, str(i+1), size=14)
    text(slide, xs[i]+0.68, 2.94, 2.1, 0.3, name, 16, NAVY, True)
    text(slide, xs[i]+0.68, 3.28, 2.1, 0.25, sub, 11, GRAY)
    if i < 3: line(slide, xs[i]+2.33, 3.25, xs[i]+2.92, 3.25, MID, 1.3, True)
text(slide, 0.82, 4.55, 11.7, 0.55, '我的理解：每篇论文都在重新选择“路径怎么走、网络学什么、迭代放在哪里”。', 20, NAVY, True)
text(slide, 0.82, 5.3, 11.5, 0.34, '重点：模型机制，而不是实验数字堆砌', 14, GRAY)
text(slide, 0.82, 6.62, 7, 0.25, 'DDPM · Song et al. · Lipman et al. · Drifting Models', 10, GRAY)
text(slide, 11.2, 6.62, 1.1, 0.25, '2026.08', 10, GRAY, align=PP_ALIGN.RIGHT)

# 2. Unified view
slide = prs.slides.add_slide(blank)
title_bar(slide, '01 / UNIFIED VIEW', '四篇论文，其实是在回答同一个问题', '如何把简单先验分布推前成真实数据分布？', BLUE)
rect(slide, 0.72, 2.03, 11.9, 1.1, LIGHT_BLUE, LIGHT_BLUE, True)
add_circle(slide, 1.05, 2.28, 0.52, NAVY, 'p₀', size=13)
text(slide, 1.68, 2.23, 1.85, 0.32, '噪声 / 先验', 16, NAVY, True)
line(slide, 3.62, 2.55, 5.05, 2.55, BLUE, 2.1, True)
text(slide, 4.0, 2.18, 1.0, 0.22, '学习映射 f', 11, BLUE, True, align=PP_ALIGN.CENTER)
add_circle(slide, 5.35, 2.28, 0.52, ORANGE, 'f', size=15)
line(slide, 6.0, 2.55, 7.58, 2.55, BLUE, 2.1, True)
text(slide, 6.24, 2.18, 1.35, 0.22, '推前 / 采样', 11, BLUE, True, align=PP_ALIGN.CENTER)
add_circle(slide, 7.9, 2.28, 0.52, GREEN, 'q', size=15)
text(slide, 8.54, 2.23, 2.8, 0.32, '逼近 p_data', 16, NAVY, True)
text(slide, 0.85, 3.62, 2.2, 0.28, '三条对照轴', 17, NAVY, True)
axes = [
    ('路径对象', '离散链 → SDE → ODE → 训练时演化', BLUE),
    ('网络预测量', '噪声 ε → 分数 s → 向量场 v → 漂移场 V', ORANGE),
    ('迭代发生处', '采样时 1000 步 → 连续求解 → 1 次前向', GREEN),
]
for i, (h, b, c) in enumerate(axes):
    y = 4.02 + i*0.72
    rect(slide, 0.84, y, 2.0, 0.48, c, c, True)
    text(slide, 0.92, y+0.08, 1.84, 0.26, h, 13, WHITE, True, align=PP_ALIGN.CENTER)
    text(slide, 3.12, y+0.08, 8.6, 0.28, b, 15, BLACK, True)
footer(slide, 2)

# 3. DDPM
slide = prs.slides.add_slide(blank)
title_bar(slide, '02 / DDPM · HO ET AL., 2020', 'DDPM：把生成拆成 1000 个小去噪步骤', '核心理解：固定前向加噪，学习反向链；ε-预测是最实用的参数化。', BLUE)
rect(slide, 0.72, 1.95, 5.1, 4.62, LIGHT, LIGHT, True)
text(slide, 0.95, 2.15, 4.6, 0.28, '模型结构', 16, BLUE, True)
add_circle(slide, 1.05, 3.0, 0.52, ORANGE, 'x₀', size=14)
text(slide, 0.82, 3.62, 0.98, 0.22, '真实图像', 11, GRAY, align=PP_ALIGN.CENTER)
line(slide, 1.65, 3.26, 2.35, 3.26, MID, 1.4, True)
text(slide, 1.77, 2.78, 1.1, 0.22, '前向加噪 q', 10, GRAY, align=PP_ALIGN.CENTER)
add_circle(slide, 2.55, 3.0, 0.52, MID, 'xₜ', BLACK, 14)
line(slide, 3.16, 3.26, 3.86, 3.26, MID, 1.4, True)
text(slide, 3.27, 2.78, 1.1, 0.22, '反向去噪 pθ', 10, GRAY, align=PP_ALIGN.CENTER)
add_circle(slide, 4.05, 3.0, 0.52, BLUE, 'xₜ₋₁', size=12)
line(slide, 4.67, 3.26, 5.23, 3.26, MID, 1.4, True)
add_circle(slide, 5.38, 3.0, 0.52, GREEN, 'x₀', size=14)
text(slide, 4.98, 3.62, 1.0, 0.22, '生成图像', 11, GRAY, align=PP_ALIGN.CENTER)
text(slide, 0.95, 4.2, 4.55, 0.42, 'xₜ = √ᾱₜ x₀ + √(1−ᾱₜ) ε', 20, NAVY, True, align=PP_ALIGN.CENTER)
text(slide, 0.97, 4.9, 4.55, 0.55, '训练时可以直接随机抽 t，一步构造任意噪声等级的 xₜ。', 13, BLACK, align=PP_ALIGN.CENTER)
text(slide, 0.95, 5.78, 4.6, 0.25, '工程模板：U-Net + 时间嵌入 + 固定方差', 12, GRAY, align=PP_ALIGN.CENTER)
bullet_box(slide, 6.15, 1.95, 6.0, 2.18, '训练学什么？', [
    '网络 εθ(xₜ,t) 预测“这一步加入的噪声”',
    'L_simple = E || ε − εθ(xₜ,t) ||²',
], BLUE, LIGHT_BLUE)
bullet_box(slide, 6.15, 4.35, 6.0, 2.22, '我对它的理解', [
    '把难问题变成大量局部去噪问题，训练稳定、实现清晰',
    '代价是采样要逐步走完反向链；它是后三篇的离散起点',
], ORANGE, LIGHT_ORANGE)
footer(slide, 3)

# 4. Score-SDE
slide = prs.slides.add_slide(blank)
title_bar(slide, '03 / SCORE-SDE · SONG ET AL., 2020', 'Score-SDE：把离散扩散提升为连续时间', '核心理解：学习每个噪声时刻的 score，就能反推 SDE；采样器可以自由替换。', CYAN)
rect(slide, 0.72, 2.0, 4.2, 4.5, LIGHT, LIGHT, True)
text(slide, 0.96, 2.2, 3.7, 0.28, '连续化之后', 16, CYAN, True)
text(slide, 1.0, 2.82, 3.6, 0.55, 'dx = f(x,t)dt + g(t)dw', 21, NAVY, True, align=PP_ALIGN.CENTER)
text(slide, 1.0, 3.58, 3.6, 0.4, '前向 SDE：逐渐破坏数据', 13, BLACK, align=PP_ALIGN.CENTER)
text(slide, 1.0, 4.2, 3.6, 0.55, '反向 SDE：\\n[f − g² ∇log pₜ]dt + g dŵ', 17, NAVY, True, align=PP_ALIGN.CENTER)
text(slide, 1.0, 5.1, 3.6, 0.4, 'score sθ(x,t) ≈ ∇ₓ log pₜ(x)', 14, CYAN, True, align=PP_ALIGN.CENTER)
text(slide, 1.0, 5.88, 3.6, 0.35, 'DDPM = VP-SDE 的离散化特例', 12, GRAY, align=PP_ALIGN.CENTER)
bullet_box(slide, 5.25, 2.0, 3.45, 2.02, '两类采样器', [
    'PC：Predictor 走 SDE，Corrector 用 Langevin 校正',
    '概率流 ODE：无随机项，可逆、可算似然',
], CYAN, LIGHT_BLUE)
bullet_box(slide, 8.95, 2.0, 3.25, 2.02, '为什么重要？', [
    '同一 score 模型可接不同求解器',
    '统一 DDPM、SMLD 与条件逆问题',
], BLUE, LIGHT)
text(slide, 5.3, 4.48, 6.5, 0.3, '统一视角', 16, NAVY, True)
line(slide, 5.4, 5.2, 7.0, 5.2, CYAN, 2.1, True)
line(slide, 7.3, 5.2, 8.9, 5.2, CYAN, 2.1, True)
line(slide, 9.2, 5.2, 11.4, 5.2, CYAN, 2.1, True)
pill(slide, 5.35, 5.62, 1.45, '数据分布', LIGHT_ORANGE, ORANGE)
pill(slide, 7.25, 5.62, 1.45, '加噪 SDE', LIGHT_BLUE, CYAN)
pill(slide, 9.15, 5.62, 1.45, 'score', LIGHT_GREEN, GREEN)
pill(slide, 11.05, 5.62, 1.1, '采样', LIGHT_RED, RED)
footer(slide, 4)

# 5. Flow Matching
slide = prs.slides.add_slide(blank)
title_bar(slide, '04 / FLOW MATCHING · LIPMAN ET AL., 2022', 'Flow Matching：不再预测 score，直接回归向量场', '核心理解：先指定一条概率路径，再学习沿路径的速度；OT 路径让轨迹更直。', ORANGE)
rect(slide, 0.72, 2.0, 5.1, 4.45, LIGHT_ORANGE, LIGHT_ORANGE, True)
text(slide, 0.98, 2.2, 4.5, 0.28, '条件流匹配（CFM）', 16, ORANGE, True)
text(slide, 1.0, 2.82, 4.55, 0.42, 'xₜ = ψₜ(x₀, x₁)', 22, NAVY, True, align=PP_ALIGN.CENTER)
text(slide, 1.0, 3.38, 4.55, 0.4, 'vθ(xₜ,t)  ≈  uₜ(xₜ | x₁)', 19, NAVY, True, align=PP_ALIGN.CENTER)
text(slide, 1.0, 4.15, 4.55, 0.44, 'L_CFM = E || vθ − uₜ ||²', 19, ORANGE, True, align=PP_ALIGN.CENTER)
text(slide, 1.0, 4.88, 4.55, 0.55, '关键定理：CFM 与 FM 的梯度相同，\\n但训练时只需逐样本计算。', 13, BLACK, align=PP_ALIGN.CENTER)
text(slide, 1.0, 5.82, 4.55, 0.28, '训练：回归速度；采样：解 ODE', 12, GRAY, align=PP_ALIGN.CENTER)
text(slide, 6.15, 2.2, 5.8, 0.28, '为什么 OT 路径更快？', 16, ORANGE, True)
line(slide, 6.35, 5.66, 11.52, 5.66, MID, 1.1)
line(slide, 6.35, 5.66, 7.0, 3.35, MID, 1.1)
line(slide, 7.0, 3.35, 8.35, 3.1, MID, 1.1)
line(slide, 8.35, 3.1, 9.55, 4.6, MID, 1.1)
line(slide, 9.55, 4.6, 11.52, 2.9, MID, 1.1)
line(slide, 6.35, 5.66, 11.52, 2.9, ORANGE, 2.6, True)
add_circle(slide, 6.1, 5.4, 0.5, BLUE, 'p₀', size=12)
add_circle(slide, 11.55, 2.65, 0.5, GREEN, 'p₁', size=12)
text(slide, 6.35, 5.97, 1.0, 0.22, '扩散路径：弯曲', 11, GRAY)
text(slide, 9.45, 2.38, 1.7, 0.22, 'OT：近似直线', 11, ORANGE, True)
rect(slide, 6.15, 6.03, 5.9, 0.52, WHITE, WHITE, True)
text(slide, 6.32, 6.13, 0.62, 0.22, '结果', 12, ORANGE, True)
text(slide, 7.05, 6.13, 4.78, 0.22, '轨迹更简单 → 训练更稳，ODE 采样 NFE 更少', 12, BLACK)
footer(slide, 5)

# 6. Drifting
slide = prs.slides.add_slide(blank)
title_bar(slide, '05 / DRIFTING MODELS · 2026', 'Drifting：把“推理时迭代”搬到“训练时演化”', '核心理解：训练过程本身让推前分布移动；平衡时漂移为零，因此推理只需一次前向。', GREEN)
rect(slide, 0.72, 2.0, 6.1, 4.48, LIGHT_GREEN, LIGHT_GREEN, True)
text(slide, 0.98, 2.2, 5.5, 0.28, '训练时的样本漂移', 16, GREEN, True)
add_circle(slide, 1.12, 3.1, 0.52, BLUE, 'ε', size=14)
line(slide, 1.72, 3.36, 2.6, 3.36, MID, 1.5, True)
add_circle(slide, 2.72, 3.1, 0.52, ORANGE, 'fθ', size=13)
line(slide, 3.33, 3.36, 4.22, 3.36, MID, 1.5, True)
add_circle(slide, 4.35, 3.1, 0.52, GREEN, 'x', size=15)
text(slide, 0.98, 3.82, 5.55, 0.32, 'x = fθ(ε)，SGD 更新 θ ⇒ x 也在移动', 15, NAVY, True, align=PP_ALIGN.CENTER)
text(slide, 0.98, 4.42, 5.55, 0.42, 'Vₚ,ᵩ(x) = 吸引真实样本 − 排斥生成样本', 15, NAVY, True, align=PP_ALIGN.CENTER)
text(slide, 0.98, 5.16, 5.55, 0.52, '反对称：Vₚ,ᵩ = −Vᵩ,ₚ\\np = q  ⇒  V = 0（达到平衡）', 14, GREEN, True, align=PP_ALIGN.CENTER)
text(slide, 0.98, 5.98, 5.55, 0.24, '训练目标：|| fθ(ε) − stopgrad(x + V) ||²', 12, GRAY, align=PP_ALIGN.CENTER)
bullet_box(slide, 7.2, 2.0, 5.0, 2.0, '工程实现', [
    '正样本：真实数据；负样本：同批生成样本',
    '漂移损失可放到多尺度特征空间',
], GREEN, LIGHT_GREEN)
bullet_box(slide, 7.2, 4.28, 5.0, 2.2, '我对它的理解', [
    '不是把多步模型蒸馏成一步，而是从训练目标上改变迭代位置',
    '优点：1-NFE；难点：漂移估计、特征空间与平衡条件',
], ORANGE, LIGHT_ORANGE)
footer(slide, 6)

# 7. Compare
slide = prs.slides.add_slide(blank)
title_bar(slide, '06 / COMPARISON', '四个模型的差异，集中在“迭代放在哪里”', '对多任务成像而言：质量、条件控制与推理预算决定选型。', BLUE)
headers = ['模型', '路径 / 状态', '网络学什么', '推理方式', '我的判断']
colx = [0.72, 2.35, 5.0, 7.8, 10.12]
colw = [1.48, 2.48, 2.6, 2.16, 2.45]
for x,w,h in zip(colx,colw,headers):
    rect(slide, x, 2.03, w, 0.48, NAVY, NAVY, True)
    text(slide, x+0.05, 2.12, w-0.1, 0.22, h, 11, WHITE, True, PP_ALIGN.CENTER)
rows = [
    ('DDPM', '离散扩散链', '噪声 ε', '多步去噪', '最适合入门与复现', BLUE, LIGHT_BLUE),
    ('Score-SDE', '连续 SDE', '分数 ∇log p', 'PC / ODE', '条件逆问题最成熟', CYAN, LIGHT_BLUE),
    ('Flow Matching', '概率路径 ODE', '向量场 vₜ', 'ODE 少步', '效率与路径可设计', ORANGE, LIGHT_ORANGE),
    ('Drifting', '训练时推前演化', '漂移场 V', '1 次前向', '实时生成最有潜力', GREEN, LIGHT_GREEN),
]
for r,(a,b,c,d,e,col,fill) in enumerate(rows):
    y = 2.56 + r*0.84
    rect(slide, colx[0], y, 11.86, 0.72, fill, WHITE, True)
    vals = [a,b,c,d,e]
    for i,(x,w,val) in enumerate(zip(colx,colw,vals)):
        text(slide, x+0.08, y+0.18, w-0.16, 0.3, val, 12 if i else 13, col if i==0 else BLACK, i==0, PP_ALIGN.CENTER if i<4 else PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE)
text(slide, 0.82, 6.17, 11.5, 0.34, '演进主线：离散去噪  →  连续 score  →  直接速度回归  →  训练时分布演化', 16, NAVY, True, align=PP_ALIGN.CENTER)
footer(slide, 7)

# 8. Closing / multi-task imaging
slide = prs.slides.add_slide(blank)
title_bar(slide, '07 / TAKEAWAY', '我的结论：先分清“要控制什么”，再决定“走哪条路径”', '四篇论文不是互相替代，而是对应不同的计算与控制需求。', BLUE)
bullet_box(slide, 0.75, 2.0, 3.75, 2.2, '如果重点是理解机制', [
    '从 DDPM 入手：前向闭式 + ε-预测 + 反向链',
    '它提供了之后三篇都能看懂的基线语言',
], BLUE, LIGHT_BLUE)
bullet_box(slide, 4.8, 2.0, 3.75, 2.2, '如果重点是多任务条件', [
    'Score-SDE 的条件 score 能统一分类、修复、超分等任务',
    '一个模型 + 不同条件项，适合做任务切换',
], CYAN, LIGHT_BLUE)
bullet_box(slide, 8.85, 2.0, 3.75, 2.2, '如果重点是推理效率', [
    'Flow Matching 先尝试 OT 路径少步采样',
    'Drifting 直接把目标推进到 1-NFE',
], GREEN, LIGHT_GREEN)
rect(slide, 0.75, 4.72, 11.85, 1.3, LIGHT, LIGHT, True)
text(slide, 1.0, 4.97, 11.35, 0.38, '一句话带走', 16, ORANGE, True, align=PP_ALIGN.CENTER)
text(slide, 1.0, 5.45, 11.35, 0.38, 'DDPM 学“怎么去噪”，Score-SDE 学“每个时刻往哪走”，Flow Matching 学“沿哪条路走”，Drifting 学“把迭代放到训练里”。', 17, NAVY, True, align=PP_ALIGN.CENTER)
text(slide, 0.78, 6.58, 11.7, 0.25, '谢谢！', 14, GRAY, align=PP_ALIGN.CENTER)
footer(slide, 8)

prs.save(OUT)
print(OUT)
