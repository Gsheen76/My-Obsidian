from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RgbColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import os

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

def add_title_slide(prs, title, subtitle):
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(12.333), Inches(1.5))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(44)
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER
    
    subtitle_box = slide.shapes.add_textbox(Inches(0.5), Inches(4.2), Inches(12.333), Inches(1))
    tf = subtitle_box.text_frame
    p = tf.paragraphs[0]
    p.text = subtitle
    p.font.size = Pt(24)
    p.alignment = PP_ALIGN.CENTER
    
    return slide

def add_content_slide(prs, title, content_list, image_path=None, image_right=False):
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.333), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(32)
    p.font.bold = True
    
    content_width = Inches(6) if image_path and image_right else Inches(12.333)
    content_left = Inches(0.5)
    
    content_box = slide.shapes.add_textbox(content_left, Inches(1.3), content_width, Inches(5.5))
    tf = content_box.text_frame
    tf.word_wrap = True
    
    for i, item in enumerate(content_list):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(20)
        p.space_after = Pt(12)
    
    if image_path and os.path.exists(image_path):
        if image_right:
            slide.shapes.add_picture(image_path, Inches(7), Inches(1.5), width=Inches(5.5))
        else:
            slide.shapes.add_picture(image_path, Inches(1.5), Inches(2), height=Inches(4))
    
    return slide

def add_image_slide(prs, title, image_path, caption=None):
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.333), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(32)
    p.font.bold = True
    
    if image_path and os.path.exists(image_path):
        slide.shapes.add_picture(image_path, Inches(2), Inches(1.3), height=Inches(5))
    
    if caption:
        caption_box = slide.shapes.add_textbox(Inches(0.5), Inches(6.5), Inches(12.333), Inches(0.5))
        tf = caption_box.text_frame
        p = tf.paragraphs[0]
        p.text = caption
        p.font.size = Pt(18)
        p.alignment = PP_ALIGN.CENTER
    
    return slide

def add_two_image_slide(prs, title, img1_path, img2_path, caption1="", caption2=""):
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.333), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(32)
    p.font.bold = True
    
    if img1_path and os.path.exists(img1_path):
        slide.shapes.add_picture(img1_path, Inches(0.5), Inches(1.5), height=Inches(4.5))
    if img2_path and os.path.exists(img2_path):
        slide.shapes.add_picture(img2_path, Inches(6.8), Inches(1.5), height=Inches(4.5))
    
    return slide

def add_three_image_slide(prs, title, img1_path, img2_path, img3_path):
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.333), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(32)
    p.font.bold = True
    
    if img1_path and os.path.exists(img1_path):
        slide.shapes.add_picture(img1_path, Inches(0.3), Inches(1.5), height=Inches(4))
    if img2_path and os.path.exists(img2_path):
        slide.shapes.add_picture(img2_path, Inches(4.5), Inches(1.5), height=Inches(4))
    if img3_path and os.path.exists(img3_path):
        slide.shapes.add_picture(img3_path, Inches(8.7), Inches(1.5), height=Inches(4))
    
    return slide

assets_dir = "assets/任务1"

add_title_slide(prs, "CT图像重建算法入门", "任务1学习汇报\n汇报人：XXX\n日期：2026年4月13日")

add_content_slide(prs, "一、Radon变换（拉东变换）", [
    "定义：",
    "• CT原始数据可近似为拉东变换",
    "• 将二维函数 f(x,y) 沿任意直线做线积分",
    "",
    "数学表达式：",
    "R(ρ,θ) = ∫ f(x,y) dl  （沿直线 xcosθ + ysinθ = ρ）",
    "",
    "参数含义：",
    "• ρ：直线到原点的距离",
    "• θ：直线的角度"
])

add_content_slide(prs, "习题1：方形图像的拉东变换", [
    "问题设置：",
    "• 边长为1的方形，内部值不均匀",
    "• 左半部分值为1，右半部分值为0.5",
    "",
    "求解结果：",
    "1. 蓝色实线（平行于边）：拉东变换值 = 1",
    "2. 红色实线（斜对角线）：拉东变换值 = 3√2/4 ≈ 1.06"
], os.path.join(assets_dir, "file-20260413111936307.png"), image_right=True)

add_content_slide(prs, "习题2：圆形图像的拉东变换", [
    "计算结果：",
    "对于半径为1的圆（圆心在原点，圆内值为1）：",
    "",
    "R(ρ,θ) = 2√(1-ρ²),  |ρ| ≤ 1",
    "",
    "特性：",
    "• 由于中心对称性，R(ρ,θ) 与 θ 无关",
    "• 只与 ρ（直线到圆心距离）有关"
])

add_image_slide(prs, "习题2：采样结果（正弦图）", 
    os.path.join(assets_dir, "file-20260413123023165.png"),
    "θ: 1°-360°（间隔1°），ρ: [-5,5]（间隔0.01）")

add_content_slide(prs, "二、图像重建的解析算法（习题3）", [
    "问题设置：",
    "• 待重建图像：2×2 维度，共4个像素",
    "• 已知：5条直线的拉东变换值",
    "",
    "求解方法：",
    "1. 线积分近似为像素值求和",
    "2. 建立5个线性方程",
    "3. 求解 x₁₁, x₁₂, x₂₁, x₂₂",
    "",
    "结论：解析方法适用于小规模图像重建"
], os.path.join(assets_dir, "file-20260413123758246.png"), image_right=True)

add_content_slide(prs, "三、傅里叶中心切片定理", [
    "定理内容：",
    "对 R(ρ,θ) 沿 ρ 方向做傅里叶变换：",
    "",
    "R̃(k,θ) = f̃(k·cosθ, k·sinθ)",
    "",
    "意义：",
    "• R̃(k,θ) 是原图傅里叶变换过原点沿 θ 方向的切片",
    "• 通过各角度切片可重建完整傅里叶变换",
    "• 再通过逆变换还原原图像"
], os.path.join(assets_dir, "file-20260413141349774.png"), image_right=True)

add_content_slide(prs, "习题4：图像重建步骤", [
    "重建流程：",
    "",
    "步骤1：计算傅里叶变换",
    "• 对 R(ρ,θ) 沿 ρ 做离散傅里叶变换得到 R̃(k,θ)",
    "",
    "步骤2：插值获得 f̃(kx,ky)",
    "• 利用傅里叶中心切片定理",
    "• 在笛卡尔网格上插值",
    "",
    "步骤3：二维傅里叶逆变换",
    "• 对 f̃(kx,ky) 做逆变换得到 f(x,y)"
])

add_three_image_slide(prs, "习题4：重建结果",
    os.path.join(assets_dir, "file-20260413145410116.png"),
    os.path.join(assets_dir, "file-20260413145410159.png"),
    os.path.join(assets_dir, "file-20260413145410181.png"))

add_content_slide(prs, "总结", [
    "1. Radon变换是CT成像的数学基础",
    "   • 将图像沿直线积分得到投影数据",
    "",
    "2. 小规模图像可用解析方法重建",
    "   • 建立线性方程组直接求解",
    "",
    "3. 大规模图像需使用滤波反投影算法",
    "   • 计算效率更高",
    "",
    "4. 傅里叶中心切片定理是高效重建的核心",
    "   • 连接投影数据与原图傅里叶变换"
])

add_title_slide(prs, "谢谢！", "")

prs.save("任务1_CT图像重建.pptx")
print("PPT已生成：任务1_CT图像重建.pptx")
